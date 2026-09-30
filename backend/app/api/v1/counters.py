import re
from datetime import date, datetime, time, timedelta

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session, joinedload

from app.core.deps import get_current_user, require
from app.db.session import get_db
from app.models.billing import Invoice, InvoiceLine
from app.models.catalog import CatalogItem, LabTestReagent
from app.models.ipd import Admission, Bed, Ward
from app.models.org import AuditLog
from app.models.patients import Patient
from app.models.users import User
from app.schemas.common import InvoiceOut, PatientOut, UserOut
from app.services.billing_service import update_line
from app.services.utils import audit
from app.schemas.counter import (
    ConvertTypeIn,
    DoctorExamIn,
    DoctorIn,
    DoctorUpdateIn,
    LabOrderCounterIn,
    LabResultIn,
    LabWalkIn,
    PharmacyCounterIn,
    RadiologyResultIn,
    ReceptionRegisterIn,
    ServiceItemIn,
    ServiceItemUpdateIn,
    StockCountIn,
    UsgOrderIn,
    XrayCounterIn,
)
from app.services.counter_service import (
    apply_stock_count,
    cashier_analytics,
    cashier_analytics_range,
    convert_patient_type,
    create_doctor,
    create_service_item,
    doctor_examine,
    get_open_invoice,
    lab_order,
    lab_walk_in,
    lab_result,
    list_doctors,
    list_service_items,
    patient_medical_history,
    search_patient_records,
    pending_prescriptions,
    pharmacy_charge,
    pharmacy_walk_in_sale,
    radiology_result,
    reception_register,
    update_doctor,
    update_service_item,
    usg_order,
    xray_charge,
)

router = APIRouter(prefix="/counter", tags=["counter"])


@router.post("/reception")
def counter_reception(data: ReceptionRegisterIn, db: Session = Depends(get_db), user: User = Depends(require("front_desk", "patients"))):
    try:
        result = reception_register(
            db,
            branch_id=data.branch_id,
            doctor_id=data.doctor_id,
            patient_id=data.patient_id,
            name=data.name,
            phone=data.phone,
            gender=data.gender,
            user_id=user.id,
            patient_type=data.patient_type,
            bed_id=data.bed_id,
            deposit=data.deposit,
            billing_mode=data.billing_mode,
            age_years=data.age_years,
            age_months=data.age_months,
            age_days=data.age_days,
            address=data.address,
            father_name=data.father_name,
            referring_doctor=data.referring_doctor,
        )
        db.commit()
        db.refresh(result["patient"])
        db.refresh(result["invoice"])
        if result.get("admission"):
            db.refresh(result["admission"])
        return {
            "patient": PatientOut.model_validate(result["patient"]),
            "patient_type": result["patient_type"],
            "invoice_id": result["invoice"].id,
            "invoice_number": result["invoice"].number,
            "token_number": result["token"].number if result["token"] else None,
            "doctor_name": result["doctor"].full_name,
            "doctor_specialty": result["doctor"].specialty or "",
            "consultation_fee": result["doctor"].consultation_fee if result["patient_type"] == "opd" else 0,
            "total": result["invoice"].total,
            "balance": result["invoice"].balance,
            "admission_id": result["admission"].id if result.get("admission") else None,
            "bed_code": result["bed"].code if result.get("bed") else None,
            "ward_name": result["ward"].name if result.get("ward") else None,
            "deposit": result["admission"].deposit if result.get("admission") else 0,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.get("/active-patients")
def active_patients(
    branch_id: int,
    q: str = "",
    doctor_id: int | None = None,
    limit: int = 50,
    with_total: bool = False,
    db: Session = Depends(get_db),
    _: User = Depends(require("pharmacy", "radiology", "front_desk", "pos", "billing.read", "patients.read")),
):
    query = (
        db.query(Invoice, Patient)
        .join(Patient, Invoice.patient_id == Patient.id)
        .options(joinedload(Invoice.doctor))
        .filter(
            Invoice.branch_id == branch_id,
            Invoice.status.in_(["draft", "open", "partial"]),
        )
    )
    if doctor_id:
        query = query.filter(Invoice.doctor_id == doctor_id)
    q = q.strip()
    if q:
        like = f"%{q}%"
        conditions = [
            Patient.name.ilike(like),
            Patient.name_mm.ilike(like),
            Patient.uhid.ilike(like),
            Patient.phone.ilike(like),
            Invoice.number.ilike(like),
            User.full_name.ilike(like),
        ]
        if q.isdigit():
            conditions.append(Patient.id == int(q))
        query = query.outerjoin(User, Invoice.doctor_id == User.id).filter(or_(*conditions))
    total = query.order_by(None).count() if with_total else None
    rows = query.order_by(Invoice.id.desc()).limit(limit).all()
    items = [
        {
            "patient_id": p.id,
            "name": p.name,
            "uhid": p.uhid,
            "phone": p.phone,
            "invoice_id": inv.id,
            "invoice_number": inv.number,
            "invoice_kind": inv.kind,
            "invoice_status": inv.status,
            "admission_id": inv.admission_id,
            "doctor_id": inv.doctor_id,
            "doctor_name": inv.doctor.full_name if inv.doctor else "",
            "total": inv.total,
            "balance": inv.balance,
        }
        for inv, p in rows
    ]
    # Existing callers (Nurse, Reception) expect a bare array — only opt-in
    # callers passing with_total=true (the "Load more" waiting-list UI) get
    # the wrapped shape with a total count to size the button/counter against.
    if with_total:
        return {"items": items, "total": total}
    return items


@router.get("/wards")
def counter_wards(branch_id: int, db: Session = Depends(get_db), _: User = Depends(require("front_desk", "pos", "patients.read"))):
    return db.query(Ward).filter(Ward.branch_id == branch_id).order_by(Ward.name).all()


@router.get("/beds")
def counter_beds(ward_id: int, available_only: bool = True, db: Session = Depends(get_db), _: User = Depends(require("front_desk", "pos", "patients.read"))):
    q = db.query(Bed).filter(Bed.ward_id == ward_id)
    if available_only:
        q = q.filter(Bed.status == "available")
    beds = q.order_by(Bed.code).all()
    return [
        {
            "id": b.id,
            "code": b.code,
            "status": b.status,
            "daily_rate": b.daily_rate,
            "hourly_rate": b.hourly_rate,
            "package_rate": b.package_rate,
        }
        for b in beds
    ]


@router.get("/patient/{patient_id}/invoice", response_model=InvoiceOut | None)
def patient_open_invoice(
    patient_id: int,
    branch_id: int,
    kind: str | None = None,
    invoice_id: int | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    if invoice_id:
        inv = (
            db.query(Invoice)
            .filter(
                Invoice.id == invoice_id,
                Invoice.patient_id == patient_id,
                Invoice.branch_id == branch_id,
                Invoice.status.in_(["draft", "open", "partial"]),
            )
            .first()
        )
    else:
        inv = get_open_invoice(db, patient_id, branch_id, kind)
    if not inv:
        return None
    inv = (
        db.query(Invoice)
        .options(joinedload(Invoice.lines), joinedload(Invoice.payments), joinedload(Invoice.doctor))
        .filter(Invoice.id == inv.id)
        .first()
    )
    inv.doctor_name = inv.doctor.full_name if inv.doctor else ""
    return inv


@router.get("/open-invoices")
def open_invoices(branch_id: int, db: Session = Depends(get_db), _: User = Depends(require("pos", "billing", "billing.read"))):
    ongoing_ipd = and_(
        Invoice.kind == "ipd",
        Invoice.admission_id.isnot(None),
        Admission.status.in_(["admitted", "transferred"]),
    )
    rows = (
        db.query(Invoice)
        .options(joinedload(Invoice.lines))
        .outerjoin(Admission, Invoice.admission_id == Admission.id)
        .filter(
            Invoice.branch_id == branch_id,
            Invoice.status.in_(["draft", "open", "partial"]),
            or_(Invoice.balance > 0.01, ongoing_ipd),
        )
        .order_by(Invoice.id.desc())
        .limit(100)
        .all()
    )
    return rows


@router.get("/xray-items")
def xray_items(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return (
        db.query(CatalogItem)
        .filter(CatalogItem.category == "radiology", CatalogItem.department != "usg", CatalogItem.is_active == True)  # noqa: E712
        .order_by(CatalogItem.price.asc())
        .all()
    )


@router.get("/pharmacy/history")
def pharmacy_history(
    branch_id: int,
    q: str = "",
    days: int = 30,
    db: Session = Depends(get_db),
    _: User = Depends(require("pharmacy", "pos")),
):
    """Dispense history for Pharmacy's own reference. Pharmacy screens do show
    unit_price/amount here (unlike most of the app being Cashier-facing for
    money) because a walk-in customer sometimes buys and pays for medicine
    directly at Pharmacy without a separate Cashier stop, so the pharmacist
    needs to quote a total."""
    cutoff = datetime.utcnow() - timedelta(days=days)
    rows = (
        db.query(InvoiceLine, Invoice, Patient)
        .join(Invoice, InvoiceLine.invoice_id == Invoice.id)
        .join(Patient, Invoice.patient_id == Patient.id)
        .filter(InvoiceLine.source == "pharmacy", Invoice.branch_id == branch_id, Invoice.created_at >= cutoff)
        .order_by(Invoice.created_at.desc())
        .limit(300)
        .all()
    )
    ql = q.strip().lower()
    out = []
    for line, inv, patient in rows:
        if ql and ql not in patient.name.lower() and ql not in (patient.uhid or "").lower() and ql not in inv.number.lower():
            continue
        out.append({
            "invoice_number": inv.number,
            "patient_name": patient.name,
            "uhid": patient.uhid,
            "item_name": line.description or (line.item.name if line.item else ""),
            "unit": line.item.unit if line.item else "",
            "qty": line.qty,
            "unit_price": line.unit_price,
            "amount": line.amount,
            "date": str(inv.created_at),
        })
    return out


@router.post("/pharmacy")
def counter_pharmacy(data: PharmacyCounterIn, db: Session = Depends(get_db), user: User = Depends(require("pharmacy", "dispense"))):
    try:
        result = pharmacy_charge(
            db,
            branch_id=data.branch_id,
            patient_id=data.patient_id,
            item_id=data.item_id,
            warehouse_id=data.warehouse_id,
            qty=data.qty,
            user_id=user.id,
            prescription_item_id=data.prescription_item_id,
            invoice_id=data.invoice_id,
        )
        db.commit()
        return {
            "ok": True,
            "invoice_id": result["invoice"].id,
            "invoice_number": result["invoice"].number,
            "total": result["invoice"].total,
            "balance": result["invoice"].balance,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.patch("/pharmacy/lines/{line_id}")
def pharmacy_line_qty(line_id: int, qty: float, db: Session = Depends(get_db), user: User = Depends(require("pharmacy", "dispense"))):
    """Qty-only correction for a pharmacist's own already-dispensed line —
    deliberately doesn't accept unit_price, so this can't become a side door
    for editing price."""
    line = db.get(InvoiceLine, line_id)
    if not line or line.source != "pharmacy":
        raise HTTPException(404)
    inv = db.get(Invoice, line.invoice_id)
    try:
        update_line(db, inv, line, qty=qty)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    audit(db, user.id, "pharmacy_line_qty", "invoice_line", str(line.id), f"qty={qty}")
    db.commit()
    return {"ok": True, "qty": line.qty}


@router.post("/pharmacy/walk-in")
def counter_pharmacy_walk_in(branch_id: int, name: str = "", phone: str = "", db: Session = Depends(get_db), user: User = Depends(require("pharmacy", "dispense"))):
    result = pharmacy_walk_in_sale(db, branch_id=branch_id, name=name, phone=phone, user_id=user.id)
    db.commit()
    inv = result["invoice"]
    patient = result["patient"]
    return {
        "patient_id": patient.id,
        "name": patient.name,
        "uhid": patient.uhid,
        "invoice_id": inv.id,
        "invoice_number": inv.number,
        "invoice_kind": inv.kind,
        "total": inv.total,
        "balance": inv.balance,
    }


@router.get("/pending-prescriptions")
def counter_pending_prescriptions(
    branch_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require("pharmacy", "dispense")),
):
    return pending_prescriptions(db, branch_id)


@router.get("/doctor/queue")
def counter_doctor_queue(
    branch_id: int,
    doctor_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require("opd", "emr")),
):
    from app.models.clinical import Visit

    rows = (
        db.query(Invoice, Patient)
        .join(Patient, Invoice.patient_id == Patient.id)
        .filter(
            Invoice.branch_id == branch_id,
            Invoice.doctor_id == doctor_id,
            Invoice.status.in_(["draft", "open", "partial"]),
        )
        .order_by(Invoice.id.desc())
        .limit(50)
        .all()
    )
    out = []
    for inv, p in rows:
        # Reception creates an empty Visit stub at registration time — only a Visit
        # with actual exam content (diagnosis/chief complaint/notes) counts as "examined".
        seen = (
            db.query(Visit)
            .filter(
                Visit.patient_id == p.id,
                Visit.doctor_id == doctor_id,
                (Visit.diagnosis != "") | (Visit.chief_complaint != "") | (Visit.notes != ""),
            )
            .order_by(Visit.id.desc())
            .first()
        )
        out.append(
            {
                "patient_id": p.id,
                "name": p.name,
                "uhid": p.uhid,
                "age_years": p.age_years,
                "gender": p.gender,
                "invoice_id": inv.id,
                "invoice_number": inv.number,
                "examined": bool(seen and seen.created_at and inv.created_at and seen.created_at >= inv.created_at),
            }
        )
    return out


@router.post("/doctor/examine")
def counter_doctor_examine(
    data: DoctorExamIn,
    db: Session = Depends(get_db),
    user: User = Depends(require("opd", "emr")),
):
    try:
        result = doctor_examine(
            db,
            branch_id=data.branch_id,
            doctor_id=data.doctor_id,
            patient_id=data.patient_id,
            chief_complaint=data.chief_complaint,
            diagnosis=data.diagnosis,
            notes=data.notes,
            follow_up_at=data.follow_up_at,
            rx_items=data.rx_items,
            user_id=user.id,
        )
        db.commit()
        return {
            "ok": True,
            "visit_id": result["visit"].id,
            "prescription_id": result["prescription"].id if result.get("prescription") else None,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.get("/lab-items")
def counter_lab_items(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    items = (
        db.query(CatalogItem)
        .filter(CatalogItem.category == "lab", CatalogItem.is_active == True)  # noqa: E712
        .order_by(CatalogItem.price.asc())
        .all()
    )
    links = (
        db.query(LabTestReagent, CatalogItem)
        .join(CatalogItem, LabTestReagent.reagent_item_id == CatalogItem.id)
        .all()
    )
    reagents: dict[int, list] = {}
    for link, reagent in links:
        reagents.setdefault(link.test_item_id, []).append({
            "item_id": reagent.id,
            "sku": reagent.sku,
            "name": reagent.name,
            "qty": link.qty,
            "unit": reagent.unit or "ea",
        })
    return [
        {
            "id": item.id,
            "sku": item.sku,
            "name": item.name,
            "name_mm": item.name_mm,
            "price": item.price,
            "category": item.category,
            "department": item.department,
            "reagents": reagents.get(item.id, []),
        }
        for item in items
    ]


@router.post("/lab")
def counter_lab(data: LabOrderCounterIn, db: Session = Depends(get_db), user: User = Depends(require("lab"))):
    try:
        result = lab_order(
            db,
            branch_id=data.branch_id,
            patient_id=data.patient_id,
            item_id=data.item_id,
            user_id=user.id,
            invoice_id=data.invoice_id,
        )
        db.commit()
        return {
            "ok": True,
            "invoice_id": result["invoice"].id,
            "invoice_number": result["invoice"].number,
            "order_id": result["order"].id,
            "item_name": result["item_name"],
            "total": result["invoice"].total,
            "balance": result["invoice"].balance,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.post("/lab/walk-in")
def counter_lab_walk_in(data: LabWalkIn, db: Session = Depends(get_db), user: User = Depends(require("lab"))):
    try:
        result = lab_walk_in(
            db,
            branch_id=data.branch_id,
            name=data.name,
            phone=data.phone,
            gender=data.gender,
            age_years=data.age_years,
            age_months=data.age_months,
            age_days=data.age_days,
            referring_doctor=data.referring_doctor,
            user_id=user.id,
        )
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    db.commit()
    inv = result["invoice"]
    patient = result["patient"]
    return {
        "patient_id": patient.id,
        "name": patient.name,
        "uhid": patient.uhid,
        "age_years": patient.age_years,
        "age_months": patient.age_months,
        "age_days": patient.age_days,
        "gender": patient.gender,
        "invoice_id": inv.id,
        "invoice_number": inv.number,
        "invoice_kind": inv.kind,
        "total": inv.total,
        "balance": inv.balance,
    }


@router.get("/lab-order/{order_id}")
def counter_lab_order_detail(order_id: int, db: Session = Depends(get_db), _: User = Depends(require("lab", "patients.read", "nursing", "ipd", "billing.read"))):
    from app.models.ancillary import LabOrder

    order = db.get(LabOrder, order_id)
    if not order:
        raise HTTPException(404, "Lab order not found")
    patient = db.get(Patient, order.patient_id)
    # The attending doctor for this specific order (carried over from the OPD
    # invoice at order time) is a far more reliable "referring doctor" than
    # Patient.referring_doctor, which is a free-text registration field that's
    # rarely filled in practice.
    doctor = db.get(User, order.doctor_id) if order.doctor_id else None
    return {
        "order_id": order.id,
        "patient_name": patient.name if patient else "",
        "uhid": patient.uhid if patient else "",
        "age_years": patient.age_years if patient else None,
        "age_months": patient.age_months if patient else None,
        "age_days": patient.age_days if patient else None,
        "gender": patient.gender if patient else "",
        "referring_doctor": doctor.full_name if doctor else (patient.referring_doctor if patient else ""),
        "tests": order.tests,
        "result": order.result,
        "status": order.status,
        "sample_id": order.sample_id,
        "created_at": order.created_at.isoformat() if order.created_at else None,
    }


@router.get("/lab/history")
def lab_history(
    branch_id: int,
    q: str = "",
    test: str = "",
    status: str = "all",
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("lab")),
):
    """Past lab orders for the lab counter's History tab. Filters combine:
    date range, status, test name, and a free-text match on patient name, ID,
    sample id, or test name."""
    from app.models.ancillary import LabOrder

    try:
        start = date.fromisoformat(date_from) if date_from else None
        end = date.fromisoformat(date_to) if date_to else None
    except ValueError as e:
        raise HTTPException(400, "date_from / date_to must be YYYY-MM-DD") from e

    query = (
        db.query(LabOrder, Patient)
        .join(Patient, LabOrder.patient_id == Patient.id)
        .filter(LabOrder.branch_id == branch_id)
    )
    if start:
        query = query.filter(LabOrder.created_at >= datetime.combine(start, time.min))
    if end:
        query = query.filter(LabOrder.created_at < datetime.combine(end + timedelta(days=1), time.min))
    if status and status != "all":
        query = query.filter(LabOrder.status == status)
    if test.strip():
        query = query.filter(LabOrder.tests.ilike(f"%{test.strip()}%"))
    ql = q.strip()
    if ql:
        like = f"%{ql}%"
        query = query.filter(or_(
            Patient.name.ilike(like),
            Patient.uhid.ilike(like),
            LabOrder.tests.ilike(like),
            LabOrder.sample_id.ilike(like),
        ))
    rows = query.order_by(LabOrder.id.desc()).limit(200).all()
    return [
        {
            "order_id": o.id,
            "patient_id": p.id,
            "patient_name": p.name,
            "uhid": p.uhid,
            "tests": o.tests,
            "result": o.result,
            "status": o.status,
            "sample_id": o.sample_id,
            "created_at": o.created_at.isoformat() if o.created_at else None,
        }
        for o, p in rows
    ]


@router.get("/lab-orders")
def counter_lab_orders(branch_id: int, status: str = "all", db: Session = Depends(get_db), _: User = Depends(require("lab"))):
    from app.models.ancillary import LabOrder

    q = db.query(LabOrder, Patient).join(Patient, LabOrder.patient_id == Patient.id).filter(LabOrder.branch_id == branch_id)
    if status != "all":
        q = q.filter(LabOrder.status == status)
    rows = q.order_by(LabOrder.id.desc()).limit(100).all()
    return [
        {
            "order_id": o.id,
            "patient_id": p.id,
            "patient_name": p.name,
            "uhid": p.uhid,
            "tests": o.tests,
            "result": o.result,
            "status": o.status,
            "sample_id": o.sample_id,
            "created_at": o.created_at.isoformat() if o.created_at else None,
        }
        for o, p in rows
    ]


@router.patch("/lab/{order_id}/result")
def counter_lab_result(order_id: int, data: LabResultIn, db: Session = Depends(get_db), user: User = Depends(require("lab"))):
    try:
        order = lab_result(db, order_id=order_id, result=data.result, user_id=user.id)
        db.commit()
        return order
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.get("/usg-items")
def counter_usg_items(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return (
        db.query(CatalogItem)
        .filter(CatalogItem.category == "radiology", CatalogItem.department == "usg", CatalogItem.is_active == True)  # noqa: E712
        .order_by(CatalogItem.price.asc())
        .all()
    )


@router.post("/usg")
def counter_usg(data: UsgOrderIn, db: Session = Depends(get_db), user: User = Depends(require("radiology"))):
    try:
        result = usg_order(db, branch_id=data.branch_id, patient_id=data.patient_id, item_id=data.item_id, user_id=user.id)
        db.commit()
        return {
            "ok": True,
            "invoice_id": result["invoice"].id,
            "invoice_number": result["invoice"].number,
            "order_id": result["order"].id,
            "item_name": result["item_name"],
            "total": result["invoice"].total,
            "balance": result["invoice"].balance,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.get("/usg/history")
def usg_history(
    branch_id: int,
    q: str = "",
    exam: str = "",
    status: str = "all",
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("radiology")),
):
    """Past USG orders. The order row only stores modality='usg'; the scan name
    is the radiology invoice line created with that order, matched in order on
    the same invoice. Filters match the X-ray history tab."""
    import json
    from collections import defaultdict

    from app.models.ancillary import RadiologyOrder

    try:
        start = date.fromisoformat(date_from) if date_from else None
        end = date.fromisoformat(date_to) if date_to else None
    except ValueError as e:
        raise HTTPException(400, "date_from / date_to must be YYYY-MM-DD") from e

    query = (
        db.query(RadiologyOrder, Patient)
        .join(Patient, RadiologyOrder.patient_id == Patient.id)
        .filter(RadiologyOrder.branch_id == branch_id, RadiologyOrder.modality == "usg")
    )
    if start:
        query = query.filter(RadiologyOrder.created_at >= datetime.combine(start, time.min))
    if end:
        query = query.filter(RadiologyOrder.created_at < datetime.combine(end + timedelta(days=1), time.min))
    if status and status != "all":
        query = query.filter(RadiologyOrder.status == status)
    rows = query.order_by(RadiologyOrder.id.desc()).limit(500).all()

    invoice_ids = {o.invoice_id for o, _p in rows if o.invoice_id}
    lines_by_invoice: dict[int, list[str]] = defaultdict(list)
    if invoice_ids:
        line_rows = (
            db.query(InvoiceLine, CatalogItem)
            .outerjoin(CatalogItem, InvoiceLine.item_id == CatalogItem.id)
            .filter(InvoiceLine.invoice_id.in_(invoice_ids), InvoiceLine.source == "radiology")
            .order_by(InvoiceLine.id.asc())
            .all()
        )
        for line, item in line_rows:
            if item is None or item.department != "usg":
                continue
            lines_by_invoice[line.invoice_id].append(line.description or item.name)

    orders_by_invoice: dict[int | None, list] = defaultdict(list)
    for pair in rows:
        orders_by_invoice[pair[0].invoice_id].append(pair)
    exam_by_order: dict[int, str] = {}
    for inv_id, pairs in orders_by_invoice.items():
        names = lines_by_invoice.get(inv_id or 0, [])
        for index, (order, _patient) in enumerate(sorted(pairs, key=lambda pair: pair[0].id)):
            exam_by_order[order.id] = names[index] if index < len(names) else ""

    def exam_label(order: RadiologyOrder) -> str:
        named = exam_by_order.get(order.id) or ""
        if named:
            return named
        try:
            parsed = json.loads(order.findings or "")
            if isinstance(parsed, dict) and parsed.get("examType"):
                return str(parsed["examType"])
        except json.JSONDecodeError:
            pass
        return "USG"

    exam_q = exam.strip().lower()
    text_q = q.strip().lower()
    out = []
    for order, patient in rows:
        label = exam_label(order)
        if exam_q and exam_q not in label.lower():
            continue
        if text_q and text_q not in patient.name.lower() and text_q not in (patient.uhid or "").lower() and text_q not in label.lower():
            continue
        out.append({
            "order_id": order.id,
            "patient_id": patient.id,
            "patient_name": patient.name,
            "uhid": patient.uhid,
            "modality": order.modality,
            "exam_name": label,
            "findings": order.findings,
            "status": order.status,
            "created_at": order.created_at.isoformat() if order.created_at else None,
        })
        if len(out) >= 200:
            break
    return out


@router.get("/xray/history")
def xray_history(
    branch_id: int,
    q: str = "",
    exam: str = "",
    status: str = "all",
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("radiology")),
):
    """Past X-ray orders for the X-ray counter's History tab. USG orders stay
    out (they use modality='usg'). Filters combine date, status, exam name,
    and a free-text match on patient name, ID, or exam."""
    from app.models.ancillary import RadiologyOrder

    try:
        start = date.fromisoformat(date_from) if date_from else None
        end = date.fromisoformat(date_to) if date_to else None
    except ValueError as e:
        raise HTTPException(400, "date_from / date_to must be YYYY-MM-DD") from e

    query = (
        db.query(RadiologyOrder, Patient, CatalogItem)
        .join(Patient, RadiologyOrder.patient_id == Patient.id)
        .outerjoin(CatalogItem, func.lower(CatalogItem.sku) == func.lower(RadiologyOrder.modality))
        .filter(RadiologyOrder.branch_id == branch_id, RadiologyOrder.modality != "usg")
    )
    if start:
        query = query.filter(RadiologyOrder.created_at >= datetime.combine(start, time.min))
    if end:
        query = query.filter(RadiologyOrder.created_at < datetime.combine(end + timedelta(days=1), time.min))
    if status and status != "all":
        query = query.filter(RadiologyOrder.status == status)
    if exam.strip():
        like = f"%{exam.strip()}%"
        query = query.filter(or_(CatalogItem.name.ilike(like), RadiologyOrder.modality.ilike(like)))
    ql = q.strip()
    if ql:
        like = f"%{ql}%"
        query = query.filter(or_(
            Patient.name.ilike(like),
            Patient.uhid.ilike(like),
            RadiologyOrder.modality.ilike(like),
            CatalogItem.name.ilike(like),
        ))
    rows = query.order_by(RadiologyOrder.id.desc()).limit(200).all()
    return [
        {
            "order_id": o.id,
            "patient_id": p.id,
            "patient_name": p.name,
            "uhid": p.uhid,
            "modality": o.modality,
            "exam_name": item.name if item else o.modality,
            "findings": o.findings,
            "status": o.status,
            "created_at": o.created_at.isoformat() if o.created_at else None,
        }
        for o, p, item in rows
    ]


@router.get("/radiology-orders")
def counter_radiology_orders(
    branch_id: int,
    modality: str = "all",
    status: str = "all",
    db: Session = Depends(get_db),
    _: User = Depends(require("radiology")),
):
    from app.models.ancillary import RadiologyOrder

    q = (
        db.query(RadiologyOrder, Patient)
        .join(Patient, RadiologyOrder.patient_id == Patient.id)
        .filter(RadiologyOrder.branch_id == branch_id)
    )
    if modality == "usg":
        q = q.filter(RadiologyOrder.modality == "usg")
    elif modality == "xray":
        q = q.filter(RadiologyOrder.modality != "usg")
    if status != "all":
        q = q.filter(RadiologyOrder.status == status)
    rows = q.order_by(RadiologyOrder.id.desc()).limit(100).all()
    return [
        {
            "order_id": o.id,
            "patient_id": p.id,
            "patient_name": p.name,
            "uhid": p.uhid,
            "modality": o.modality,
            "findings": o.findings,
            "status": o.status,
            "created_at": o.created_at.isoformat() if o.created_at else None,
        }
        for o, p in rows
    ]


@router.get("/radiology/{order_id}")
def counter_radiology_order_detail(order_id: int, db: Session = Depends(get_db), _: User = Depends(require("radiology", "patients.read", "nursing", "ipd", "billing.read"))):
    from app.models.ancillary import RadiologyOrder

    order = db.get(RadiologyOrder, order_id)
    if not order:
        raise HTTPException(404, "Radiology order not found")
    patient = db.get(Patient, order.patient_id)
    # See counter_lab_order_detail: the order's own doctor_id (the OPD
    # invoice's attending doctor at order time) is far more reliable than
    # Patient.referring_doctor, which is rarely filled in practice.
    doctor = db.get(User, order.doctor_id) if order.doctor_id else None
    return {
        "order_id": order.id,
        "patient_name": patient.name if patient else "",
        "uhid": patient.uhid if patient else "",
        "age_years": patient.age_years if patient else None,
        "age_months": patient.age_months if patient else None,
        "age_days": patient.age_days if patient else None,
        "gender": patient.gender if patient else "",
        "referring_doctor": doctor.full_name if doctor else (patient.referring_doctor if patient else ""),
        "modality": order.modality,
        "findings": order.findings,
        "status": order.status,
        "created_at": order.created_at.isoformat() if order.created_at else None,
    }


@router.patch("/radiology/{order_id}/result")
def counter_radiology_result(order_id: int, data: RadiologyResultIn, db: Session = Depends(get_db), user: User = Depends(require("radiology"))):
    try:
        order = radiology_result(db, order_id=order_id, findings=data.findings, user_id=user.id)
        db.commit()
        return order
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.post("/convert-type")
def counter_convert_type(data: ConvertTypeIn, db: Session = Depends(get_db), user: User = Depends(require("front_desk", "patients"))):
    try:
        result = convert_patient_type(
            db,
            branch_id=data.branch_id,
            patient_id=data.patient_id,
            target_type=data.target_type,
            doctor_id=data.doctor_id,
            bed_id=data.bed_id,
            deposit=data.deposit,
            billing_mode=data.billing_mode,
            user_id=user.id,
        )
        db.commit()
        return {
            "ok": True,
            "patient_type": result["patient_type"],
            "invoice_id": result["invoice"].id,
            "invoice_number": result["invoice"].number,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.post("/xray")
def counter_xray(data: XrayCounterIn, db: Session = Depends(get_db), user: User = Depends(require("radiology"))):
    try:
        result = xray_charge(
            db,
            branch_id=data.branch_id,
            patient_id=data.patient_id,
            item_id=data.item_id,
            user_id=user.id,
        )
        db.commit()
        return {
            "ok": True,
            "invoice_id": result["invoice"].id,
            "invoice_number": result["invoice"].number,
            "order_id": result["order"].id,
            "item_name": result["item_name"],
            "total": result["invoice"].total,
            "balance": result["invoice"].balance,
        }
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.post("/stock-count")
def stock_count(data: StockCountIn, db: Session = Depends(get_db), user: User = Depends(require("pos", "inventory"))):
    try:
        result = apply_stock_count(db, data.item_id, data.warehouse_id, data.counted_qty, user.id, data.notes)
        db.commit()
        return result
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


def _bill_history_rows(db: Session, branch_id: int, q: str, status: str, days: int) -> list[dict]:
    since = datetime.utcnow() - timedelta(days=max(days, 1))
    query = (
        db.query(Invoice, Patient)
        .outerjoin(Patient, Invoice.patient_id == Patient.id)
        .options(joinedload(Invoice.payments))
        .filter(Invoice.branch_id == branch_id, Invoice.created_at >= since)
    )
    if status and status != "all":
        query = query.filter(Invoice.status == status)
    for term in re.split(r"[,\s]+", q.strip()):
        if not term:
            continue
        like = f"%{term}%"
        query = query.filter(
            (Patient.name.ilike(like))
            | (Patient.father_name.ilike(like))
            | (Patient.uhid.ilike(like))
            | (Patient.phone.ilike(like))
            | (Invoice.number.ilike(like))
        )
    rows = query.order_by(Invoice.id.desc()).limit(200).all()
    out = []
    for inv, patient in rows:
        methods = ", ".join(sorted({p.method.upper() for p in inv.payments if p.method != "refund"}))
        out.append(
            {
                "id": inv.id,
                "number": inv.number,
                "patient_id": inv.patient_id,
                "patient_name": patient.name if patient else "—",
                "uhid": patient.uhid if patient else "—",
                "phone": patient.phone if patient else "",
                "status": inv.status,
                "total": inv.total,
                "paid": inv.paid,
                "balance": inv.balance,
                "created_at": inv.created_at.isoformat() if inv.created_at else None,
                "payment_methods": methods or "—",
                "invoice_kind": inv.kind,
            }
        )
    return out


@router.get("/patient-records/search")
def counter_patient_records_search(
    branch_id: int,
    date_from: str = "",
    date_to: str = "",
    kind: str = "",
    source: str = "",
    q: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("front_desk", "patients.read", "patients")),
):
    from datetime import date

    try:
        start = date.fromisoformat(date_from) if date_from else None
        end = date.fromisoformat(date_to) if date_to else None
    except ValueError as e:
        raise HTTPException(400, "date_from / date_to must be YYYY-MM-DD") from e
    try:
        return search_patient_records(
            db,
            branch_id=branch_id,
            date_from=start,
            date_to=end,
            kind=kind,
            source=source,
            q=q,
        )
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.get("/patient-medical-history")
def counter_patient_medical_history(
    patient_id: int,
    branch_id: int | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(require("front_desk", "patients.read", "patients")),
):
    try:
        result = patient_medical_history(db, patient_id=patient_id, branch_id=branch_id)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    return {
        "patient": PatientOut.model_validate(result["patient"]),
        "visits": result["visits"],
        "clinical_notes": result["clinical_notes"],
    }


@router.get("/bill-history")
def bill_history(
    branch_id: int,
    q: str = "",
    status: str = "all",
    days: int = 30,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing.read", "billing", "front_desk", "patients.read")),
):
    return _bill_history_rows(db, branch_id, q, status, days)


@router.get("/bill-history/export")
def bill_history_export(
    branch_id: int,
    q: str = "",
    status: str = "all",
    days: int = 30,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing.read", "billing")),
):
    import openpyxl
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    rows = _bill_history_rows(db, branch_id, q, status, days)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bill History"
    ws.append(["Bill No.", "Type", "Patient", "UHID", "Phone", "Status", "Total", "Paid", "Balance", "Paid Via", "Date"])
    for r in rows:
        ws.append([
            r["number"], r["invoice_kind"], r["patient_name"], r["uhid"], r["phone"],
            r["status"], r["total"], r["paid"], r["balance"], r["payment_methods"],
            r["created_at"][:19].replace("T", " ") if r["created_at"] else "",
        ])
    ws.append([])
    ws.append(["", "", "", "", "", "Totals", sum(r["total"] for r in rows), sum(r["paid"] for r in rows), sum(r["balance"] for r in rows)])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=bill-history.xlsx"},
    )


@router.get("/doctors", response_model=list[UserOut])
def counter_doctors(
    specialty: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing", "front_desk", "patients.read")),
):
    return list_doctors(db, specialty=specialty)


@router.post("/doctors", response_model=UserOut)
def add_doctor(
    data: DoctorIn,
    branch_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require("pos", "billing")),
):
    try:
        doctor = create_doctor(
            db,
            full_name=data.full_name,
            consultation_fee=data.consultation_fee,
            branch_id=branch_id,
            specialty=data.specialty,
            user_id=user.id,
        )
        db.commit()
        db.refresh(doctor)
        return doctor
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.patch("/doctors/{doctor_id}", response_model=UserOut)
def patch_doctor(
    doctor_id: int,
    data: DoctorUpdateIn,
    db: Session = Depends(get_db),
    user: User = Depends(require("pos", "billing")),
):
    try:
        doctor = update_doctor(
            db,
            doctor_id,
            full_name=data.full_name,
            consultation_fee=data.consultation_fee,
            specialty=data.specialty,
            user_id=user.id,
        )
        db.commit()
        db.refresh(doctor)
        return doctor
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.get("/services")
def counter_list_services(
    department: str,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing")),
):
    try:
        return list_service_items(db, department)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.post("/services")
def counter_create_service(
    data: ServiceItemIn,
    db: Session = Depends(get_db),
    user: User = Depends(require("pos", "billing")),
):
    try:
        item = create_service_item(db, department=data.department, name=data.name, price=data.price, user_id=user.id)
        db.commit()
        db.refresh(item)
        return item
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


@router.patch("/services/{item_id}")
def counter_update_service(
    item_id: int,
    data: ServiceItemUpdateIn,
    db: Session = Depends(get_db),
    user: User = Depends(require("pos", "billing")),
):
    try:
        item = update_service_item(
            db,
            item_id,
            name=data.name,
            price=data.price,
            is_active=data.is_active,
            user_id=user.id,
        )
        db.commit()
        db.refresh(item)
        return item
    except ValueError as e:
        raise HTTPException(400, str(e)) from e


def _resolve_analytics(db: Session, branch_id: int, period: str, date_from: str | None, date_to: str | None) -> dict:
    if date_from and date_to:
        try:
            start = datetime.fromisoformat(date_from)
            end = datetime.fromisoformat(date_to).replace(hour=23, minute=59, second=59)
        except ValueError as e:
            raise HTTPException(400, "date_from/date_to must be YYYY-MM-DD") from e
        if end < start:
            raise HTTPException(400, "date_to must be on or after date_from")
        return cashier_analytics_range(db, branch_id, start, end)
    if period not in {"day", "month", "3m", "year"}:
        raise HTTPException(400, "period must be day, month, 3m, or year")
    return cashier_analytics(db, branch_id, period)


@router.get("/analytics")
def counter_analytics(
    branch_id: int,
    period: str = "month",
    date_from: str | None = None,
    date_to: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing", "accounts", "analytics", "reports")),
):
    return _resolve_analytics(db, branch_id, period, date_from, date_to)


@router.get("/analytics/export")
def counter_analytics_export(
    branch_id: int,
    period: str = "month",
    date_from: str | None = None,
    date_to: str | None = None,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing", "accounts", "analytics", "reports")),
):
    import openpyxl
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    a = _resolve_analytics(db, branch_id, period, date_from, date_to)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Period", a.get("period_label", "")])
    ws.append(["From", str(a.get("from_date", ""))[:19].replace("T", " ")])
    ws.append(["To", str(a.get("to_date", ""))[:19].replace("T", " ")])
    ws.append([])
    ws.append(["Cash In", a.get("total_collected", 0)])
    ws.append(["Today's Billing", a.get("total_billed", 0)])
    ws.append(["Expenses", a.get("total_expenses", 0)])
    ws.append(["Net", a.get("net", 0)])
    ws.append(["Outstanding", a.get("outstanding", 0)])
    ws.append([])
    ws.append(["Total Bills", a.get("total_bills", 0)])
    ws.append(["Paid Bills", a.get("paid_bills", 0)])
    ws.append(["Avg per Paid Bill", a.get("avg_bill", 0)])

    ws2 = wb.create_sheet("Payment Methods")
    ws2.append(["Method", "Amount"])
    for p in a.get("payment_methods", []) or []:
        ws2.append([p.get("method", ""), p.get("amount", 0)])

    ws3 = wb.create_sheet("By Department")
    ws3.append(["Department", "Amount"])
    for s in a.get("by_source", []) or []:
        ws3.append([s.get("source", ""), s.get("amount", 0)])

    ws4 = wb.create_sheet("Doctor Income")
    ws4.append(["Doctor", "Bills", "Billed", "Fee Income"])
    for d in a.get("by_doctor", []) or []:
        ws4.append([d.get("doctor", ""), d.get("bills", 0), d.get("billed", 0), d.get("fee_income", 0)])

    ws5 = wb.create_sheet("Expenses")
    ws5.append(["Date", "Category", "Paid From", "Notes", "Amount"])
    for e in a.get("expense_details", []) or []:
        ws5.append([
            str(e.get("created_at", ""))[:19].replace("T", " "),
            e.get("category", ""),
            e.get("paid_from", ""),
            e.get("notes", ""),
            e.get("amount", 0),
        ])
    ws5.append([])
    ws5.append(["", "", "", "Total", a.get("total_expenses", 0)])

    ws6 = wb.create_sheet("Trend")
    ws6.append(["Label", "Amount"])
    for t in a.get("trend", []) or []:
        ws6.append([t.get("label", ""), t.get("amount", 0)])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=analytics-report.xlsx"},
    )


@router.get("/activity")
def activity_log(
    limit: int = 100,
    db: Session = Depends(get_db),
    _: User = Depends(require("pos", "billing.read", "front_desk", "pharmacy", "radiology", "patients.read")),
):
    rows = (
        db.query(AuditLog, User)
        .outerjoin(User, AuditLog.user_id == User.id)
        .order_by(AuditLog.id.desc())
        .limit(min(limit, 300))
        .all()
    )
    return [
        {
            "id": log.id,
            "action": log.action,
            "entity": log.entity,
            "entity_id": log.entity_id,
            "detail": log.detail,
            "user_name": user.full_name if user else "—",
            "created_at": log.created_at.isoformat() if log.created_at else None,
        }
        for log, user in rows
    ]
