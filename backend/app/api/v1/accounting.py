from datetime import date

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from io import BytesIO
from sqlalchemy.orm import Session

from app.core.deps import require
from app.db.session import get_db
from app.models.accounting import Expense, JournalEntry, LedgerAccount, PettyCash
from app.models.billing import Payment
from app.models.patients import CorporateAccount
from app.models.users import User
from app.schemas.actions import ExpenseIn
from app.services.accounting_service import post_expense

router = APIRouter(prefix="/accounts", tags=["accounts"])


@router.get("/ledger")
def ledger(db: Session = Depends(get_db), _: User = Depends(require("accounts", "reports"))):
    return db.query(LedgerAccount).order_by(LedgerAccount.code).all()


@router.get("/journal")
def journal(db: Session = Depends(get_db), _: User = Depends(require("accounts", "reports"))):
    return db.query(JournalEntry).order_by(JournalEntry.id.desc()).limit(100).all()


@router.get("/cash-book")
def cash_book(db: Session = Depends(get_db), _: User = Depends(require("accounts", "reports"))):
    return db.query(Payment).order_by(Payment.id.desc()).limit(200).all()


@router.get("/ar")
def accounts_receivable(db: Session = Depends(get_db), _: User = Depends(require("accounts", "reports"))):
    from app.models.billing import Invoice

    return db.query(Invoice).filter(Invoice.balance > 0).order_by(Invoice.id.desc()).limit(100).all()


@router.get("/ap")
def accounts_payable(db: Session = Depends(get_db), _: User = Depends(require("accounts", "reports"))):
    from app.models.inventory import Supplier

    return db.query(Supplier).filter(Supplier.balance > 0).all()


@router.get("/corporate")
def corporate_accounts(db: Session = Depends(get_db), _: User = Depends(require("accounts", "billing.read"))):
    return db.query(CorporateAccount).all()


@router.post("/expenses")
def create_expense(data: ExpenseIn, db: Session = Depends(get_db), user: User = Depends(require("accounts", "pos"))):
    from app.models.billing import CashierShift

    shift = db.query(CashierShift).filter(CashierShift.user_id == user.id, CashierShift.status == "open").first()
    exp = Expense(
        branch_id=data.branch_id,
        shift_id=shift.id if shift else None,
        category=data.category,
        name=data.name,
        amount=data.amount,
        paid_from=data.paid_from,
        paid_by=data.paid_by,
        notes=data.notes,
    )
    db.add(exp)
    post_expense(db, data.category, data.amount, data.paid_from)
    # The petty-cash balance tracks the physical cash box specifically — an
    # expense paid via KPay/Wave moves money out of that e-wallet, not the
    # till, so it shouldn't touch this balance (it already posts to its own
    # ledger account — 1010/1011 — via post_expense() above).
    petty = db.query(PettyCash).filter(PettyCash.branch_id == data.branch_id).first()
    if petty and data.paid_from == "cash":
        petty.balance -= data.amount
    db.commit()
    db.refresh(exp)
    return exp


@router.get("/expenses")
def list_expenses(
    branch_id: int,
    date_from: str | None = None,
    date_to: str | None = None,
    category: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("accounts", "pos")),
):
    q = db.query(Expense).filter(Expense.branch_id == branch_id)
    if date_from:
        q = q.filter(Expense.expense_date >= date_from)
    if date_to:
        q = q.filter(Expense.expense_date <= date_to)
    if category.strip():
        q = q.filter(Expense.category == category.strip())
    rows = q.order_by(Expense.expense_date.desc(), Expense.id.desc()).limit(500).all()
    return {
        "rows": rows,
        "total": sum(r.amount for r in rows),
        "count": len(rows),
    }


@router.get("/expenses/categories")
def expense_categories(branch_id: int, db: Session = Depends(get_db), _: User = Depends(require("accounts", "pos"))):
    rows = (
        db.query(Expense.category)
        .filter(Expense.branch_id == branch_id, Expense.category != "")
        .distinct()
        .order_by(Expense.category)
        .all()
    )
    return [r[0] for r in rows]


@router.get("/expenses/export")
def export_expenses(
    branch_id: int,
    date_from: str | None = None,
    date_to: str | None = None,
    category: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("accounts", "pos")),
):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    q = db.query(Expense).filter(Expense.branch_id == branch_id)
    if date_from:
        q = q.filter(Expense.expense_date >= date_from)
    if date_to:
        q = q.filter(Expense.expense_date <= date_to)
    if category.strip():
        q = q.filter(Expense.category == category.strip())
    rows = q.order_by(Expense.category, Expense.expense_date.desc(), Expense.id.desc()).limit(2000).all()

    by_category: dict[str, list[Expense]] = {}
    for r in rows:
        by_category.setdefault(r.category or "Uncategorized", []).append(r)
    sorted_categories = sorted(by_category.items(), key=lambda kv: -sum(i.amount for i in kv[1]))
    grand_total = sum(r.amount for r in rows)

    thin = Side(style="thin", color="D0D5DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="005293")
    header_font = Font(bold=True, color="FFFFFF")
    subtotal_fill = PatternFill("solid", fgColor="EAF2FA")
    total_fill = PatternFill("solid", fgColor="FDECEC")
    bold = Font(bold=True)
    money_fmt = "#,##0"
    COLS = ["No", "Date", "Category", "Name", "Paid From", "Paid By", "Amount"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"

    ws.append(["Expense History Report"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Period: {date_from or 'all'} — {date_to or 'all'}" + (f"  ·  Category: {category}" if category.strip() else "")])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([None] * len(COLS))

    ws.append(COLS)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = header_font
        c.fill = header_fill
        c.border = border
        c.alignment = Alignment(horizontal="center")

    no = 0
    for cat, items in sorted_categories:
        for r in items:
            no += 1
            ws.append([no, str(r.expense_date), cat, r.name, r.paid_from, r.paid_by, r.amount])
            row = ws.max_row
            for c in ws[row]:
                c.border = border
            ws.cell(row, 7).number_format = money_fmt

        ws.append(["", "", "", "", "", f"{cat} Subtotal", sum(i.amount for i in items)])
        row = ws.max_row
        for c in ws[row]:
            c.border = border
            c.fill = subtotal_fill
            c.font = bold
        ws.cell(row, 7).number_format = money_fmt

    ws.append(["", "", "", "", "", "GRAND TOTAL", grand_total])
    row = ws.max_row
    for c in ws[row]:
        c.border = border
        c.fill = total_fill
        c.font = bold
    ws.cell(row, 7).number_format = money_fmt

    widths = [6, 14, 18, 20, 12, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(header_row + 1, 1)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=expense-history.xlsx"},
    )


@router.get("/expenses/export/pdf")
def export_expenses_pdf(
    branch_id: int,
    date_from: str | None = None,
    date_to: str | None = None,
    category: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require("accounts", "pos")),
):
    from app.services.pdf_service import build_table_pdf

    q = db.query(Expense).filter(Expense.branch_id == branch_id)
    if date_from:
        q = q.filter(Expense.expense_date >= date_from)
    if date_to:
        q = q.filter(Expense.expense_date <= date_to)
    if category.strip():
        q = q.filter(Expense.category == category.strip())
    rows = q.order_by(Expense.category, Expense.expense_date.desc(), Expense.id.desc()).limit(2000).all()
    grand_total = sum(r.amount for r in rows)

    table_rows = [[r.expense_date, r.category or "Uncategorized", r.name, r.paid_from, r.paid_by, f"{r.amount:,.0f}"] for r in rows]
    table_rows.append(["", "", "", "", "GRAND TOTAL", f"{grand_total:,.0f}"])

    buf = build_table_pdf(
        title="Expense History Report",
        meta=[("Period", f"{date_from or 'all'} — {date_to or 'all'}"), ("Category", category or "All")],
        columns=["Date", "Category", "Name", "Paid From", "Paid By", "Amount"],
        rows=table_rows,
        landscape_mode=True,
    )
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="expense-history.pdf"'},
    )


@router.get("/petty-cash")
def petty(branch_id: int, db: Session = Depends(get_db), _: User = Depends(require("accounts", "pos"))):
    return db.query(PettyCash).filter(PettyCash.branch_id == branch_id).first()


@router.get("/reports/pl")
def pl(db: Session = Depends(get_db), _: User = Depends(require("reports", "accounts"))):
    rev = db.query(LedgerAccount).filter(LedgerAccount.code == "4000").first()
    exp = db.query(LedgerAccount).filter(LedgerAccount.code == "5000").first()
    return {"revenue": rev.balance if rev else 0, "expense": exp.balance if exp else 0, "profit": (rev.balance if rev else 0) - (exp.balance if exp else 0)}


@router.get("/export/excel")
def export_excel(db: Session = Depends(get_db), _: User = Depends(require("reports", "accounts"))):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["Code", "Name", "Kind", "Balance"])
    for acc in db.query(LedgerAccount).all():
        ws.append([acc.code, acc.name, acc.kind, acc.balance])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=ledger.xlsx"})
