from io import BytesIO

import openpyxl


def test_create_expense_returns_full_saved_record(client, cashier_headers):
    """Regression test: create_expense used to return `{}` because it never
    called db.refresh() after commit expired the ORM object's attributes."""
    r = client.post(
        "/api/v1/accounts/expenses",
        json={"branch_id": 1, "category": "test-suite-utilities", "name": "September electricity bill", "amount": 12345, "paid_from": "petty", "paid_by": "Cashier One", "notes": "pytest"},
        headers=cashier_headers,
    )
    assert r.status_code == 200
    data = r.json()
    assert data["id"] is not None
    assert data["category"] == "test-suite-utilities"
    assert data["name"] == "September electricity bill"
    assert data["amount"] == 12345
    assert data["paid_by"] == "Cashier One"
    assert data["notes"] == "pytest"


def test_expense_categories_lists_distinct_previously_used_categories(client, cashier_headers):
    client.post(
        "/api/v1/accounts/expenses",
        json={"branch_id": 1, "category": "test-suite-repeat-category", "amount": 500, "paid_from": "petty"},
        headers=cashier_headers,
    )
    client.post(
        "/api/v1/accounts/expenses",
        json={"branch_id": 1, "category": "test-suite-repeat-category", "amount": 700, "paid_from": "petty"},
        headers=cashier_headers,
    )

    r = client.get("/api/v1/accounts/expenses/categories", params={"branch_id": 1}, headers=cashier_headers)
    assert r.status_code == 200
    categories = r.json()
    assert categories.count("test-suite-repeat-category") == 1


def test_expenses_list_filters_by_category_and_date_range(client, cashier_headers):
    client.post(
        "/api/v1/accounts/expenses",
        json={"branch_id": 1, "category": "test-suite-filter-me", "amount": 999, "paid_from": "petty"},
        headers=cashier_headers,
    )

    r = client.get("/api/v1/accounts/expenses", params={"branch_id": 1, "category": "test-suite-filter-me"}, headers=cashier_headers)
    assert r.status_code == 200
    data = r.json()
    assert data["count"] == 1
    assert data["total"] == 999
    assert data["rows"][0]["category"] == "test-suite-filter-me"

    far_future = "2999-01-01"
    r = client.get(
        "/api/v1/accounts/expenses",
        params={"branch_id": 1, "category": "test-suite-filter-me", "date_from": far_future},
        headers=cashier_headers,
    )
    assert r.json()["count"] == 0


def test_expenses_export_is_a_valid_grouped_xlsx(client, cashier_headers):
    client.post(
        "/api/v1/accounts/expenses",
        json={"branch_id": 1, "category": "test-suite-export-a", "amount": 1000, "paid_from": "petty"},
        headers=cashier_headers,
    )
    client.post(
        "/api/v1/accounts/expenses",
        json={"branch_id": 1, "category": "test-suite-export-a", "amount": 2000, "paid_from": "petty"},
        headers=cashier_headers,
    )

    r = client.get(
        "/api/v1/accounts/expenses/export",
        params={"branch_id": 1, "category": "test-suite-export-a"},
        headers=cashier_headers,
    )
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]

    wb = openpyxl.load_workbook(BytesIO(r.content))
    ws = wb.active
    values = [row[0] for row in ws.iter_rows(values_only=True)]
    amounts = [row[6] for row in ws.iter_rows(values_only=True) if row[6] is not None]
    assert "test-suite-export-a Subtotal" not in values  # subtotal label is in column F, not A
    assert 3000 in amounts  # subtotal or grand total for the two rows above
